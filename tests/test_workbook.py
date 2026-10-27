"""The loader must return exactly what is in the workbook.

These tests read cells by fixed address (independently of the label-based
loader) and compare the two.
"""

import openpyxl
import pytest

from carbon_explorer.data import HORIZON, WORKBOOK_PATH, load_workbook


@pytest.fixture(scope="module")
def wb():
    return load_workbook()


@pytest.fixture(scope="module")
def raw():
    return openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)


def test_families_in_workbook_order(wb, raw):
    sheet = raw["Portfolio"]
    assert wb.families == tuple(sheet.cell(r, 1).value for r in range(6, 19))
    assert len(wb.families) == 13


def test_quantities_match_cells(wb, raw):
    sheet = raw["Portfolio"]
    for col in range(2, 17):  # B..P = 2021..2035
        year = sheet.cell(5, col).value
        for r, family in enumerate(wb.families, start=6):
            assert wb.quantities.at[family, year] == sheet.cell(r, col).value
        assert wb.quantity_subtotals[year] == sheet.cell(19, col).value


def test_horizon_years_present(wb):
    assert HORIZON == tuple(range(2026, 2036))
    assert all(y in wb.years for y in HORIZON)


def test_emission_factors_match_cells(wb, raw):
    sheet = raw["Est.Emissions"]
    for r, family in enumerate(wb.families, start=5):
        assert sheet.cell(r, 1).value == family
        assert wb.emissions.at[family, "per_flight"] == sheet.cell(r, 2).value
        assert wb.emissions.at[family, "flights"] == sheet.cell(r, 3).value
        assert wb.emissions.at[family, "annual"] == sheet.cell(r, 4).value


def test_economics_match_cells(wb, raw):
    sheet = raw["Portfolio"]
    for r, family in enumerate(wb.families, start=6):
        assert sheet.cell(r, 20).value == family
        for key, col in (("income", 21), ("depreciation", 22), ("interest", 23), ("profit", 24)):
            assert wb.economics.at[family, key] == pytest.approx(sheet.cell(r, col).value)
    assert wb.economics_year == "2026"


def test_projects_match_cells_and_skip_placeholder(wb, raw):
    sheet = raw["CO2.Credits"]
    expected = [
        (sheet.cell(r, 2).value, round(sheet.cell(r, 6).value, 2), sheet.cell(r, 3).hyperlink.target)
        for r in range(4, 10)
    ]
    assert [(p.name, p.price, p.url) for p in wb.projects] == expected
    assert all(not p.name.lower().startswith("select") for p in wb.projects)
    prices = {p.name: p.price for p in wb.projects}
    assert prices["Aranya Reforestation"] == 21.80
    assert prices["Sabah INFAPRO Rainforest IFM"] == 28.20
    assert prices["Bondhu Chula Stoves"] == 14.80


def test_project_without_country_is_none(wb):
    assert wb.project("Accelerate CDR").country is None
    assert wb.project("Aranya Reforestation").country == "India"


def test_notes_flag_url_mismatch_and_placeholder(wb):
    text = " ".join(wb.notes)
    assert "placeholder" in text
    assert "Project URL" in text
