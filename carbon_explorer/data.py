"""Load the RRPF carbon workbook into typed, read-only tables.

The workbook is the single source of truth. Tables are located by their header
labels rather than fixed cell addresses, so a revised workbook that keeps the
same headings loads without code changes. Values are read as Excel last
calculated them (cached formula results).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd

WORKBOOK_PATH = (
    Path(__file__).resolve().parent.parent / "data" / "Carbon_Emissions_Workbook_rev.xlsx"
)

# The modelling horizon used throughout the app.
HORIZON: tuple[int, ...] = tuple(range(2026, 2036))

PORTFOLIO_SHEET = "Portfolio"
EMISSIONS_SHEET = "Est.Emissions"
CREDITS_SHEET = "CO2.Credits"


class WorkbookError(ValueError):
    """Raised when the workbook does not have the expected structure."""


@dataclass(frozen=True)
class Project:
    name: str
    country: str | None
    kind: str
    price: float  # USD per tCO2e
    url: str | None

    @property
    def label(self) -> str:
        where = self.country or "Country not stated"
        return f"{self.name} · {self.kind} · {where} · ${self.price:,.2f}/tCO₂e"


@dataclass(frozen=True)
class Workbook:
    families: tuple[str, ...]
    # Engine quantity forecast: index = engine family, columns = years (int).
    quantities: pd.DataFrame
    # The workbook's own "Sub-Totals" row, by year.
    quantity_subtotals: dict[int, float]
    # Index = engine family; columns income, depreciation, interest (monthly,
    # per engine) and profit (average profit per engine per year).
    economics: pd.DataFrame
    economics_title: str
    # Index = engine family; columns per_flight, flights, annual (tCO2e).
    emissions: pd.DataFrame
    projects: tuple[Project, ...]
    # Observations about the source data made while loading.
    notes: tuple[str, ...] = field(default_factory=tuple)
    source_name: str = WORKBOOK_PATH.name

    @property
    def years(self) -> tuple[int, ...]:
        return tuple(int(y) for y in self.quantities.columns)

    @property
    def economics_year(self) -> str | None:
        """The year named in the economics table title, e.g. '2026'."""
        match = re.search(r"\b(?:19|20)\d{2}\b", self.economics_title)
        return match.group(0) if match else None

    def quantities_for(self, year: int) -> dict[str, int]:
        if year not in self.quantities.columns:
            raise KeyError(f"{year} is not a year in the workbook forecast")
        return {f: int(self.quantities.at[f, year]) for f in self.families}

    def annual_tco2e(self, family: str) -> float:
        return float(self.emissions.at[family, "annual"])

    def profit(self, family: str) -> float:
        return float(self.economics.at[family, "profit"])

    def project(self, name: str) -> Project:
        for p in self.projects:
            if p.name == name:
                return p
        raise KeyError(f"Unknown carbon project: {name}")


# ---------------------------------------------------------------------------
# Sheet parsing helpers


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _find_cell(ws, predicate, *, start_row: int = 1):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            if predicate(_text(cell.value)):
                return cell
    return None


def _require(cell, what: str, sheet: str):
    if cell is None:
        raise WorkbookError(f"Could not find {what} on sheet '{sheet}'")
    return cell


def _header_columns(ws, row: int, first_col: int, last_col: int) -> dict[str, int]:
    return {
        _text(ws.cell(row, col).value): col
        for col in range(first_col, last_col + 1)
        if _text(ws.cell(row, col).value)
    }


def _column_matching(headers: dict[str, int], *keywords: str, sheet: str) -> int:
    for header, col in headers.items():
        lowered = header.lower()
        if all(k.lower() in lowered for k in keywords):
            return col
    raise WorkbookError(f"No column matching {keywords} on sheet '{sheet}'")


def _number(value, where: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise WorkbookError(f"Expected a number at {where}, found {value!r}")
    return float(value)


def _parse_portfolio(ws):
    sheet = ws.title
    header = _require(
        _find_cell(ws, lambda v: v == "Type"), "the quantity table header 'Type'", sheet
    )
    hrow = header.row

    # Year columns run right of the "Type" header until the first non-year cell.
    year_cols: dict[int, int] = {}
    col = header.column + 1
    while isinstance(ws.cell(hrow, col).value, (int, float)):
        year_cols[int(ws.cell(hrow, col).value)] = col
        col += 1
    if not year_cols:
        raise WorkbookError(f"No year columns found on sheet '{sheet}'")
    missing = [y for y in HORIZON if y not in year_cols]
    if missing:
        raise WorkbookError(f"Sheet '{sheet}' is missing forecast years {missing}")

    families: list[str] = []
    rows: dict[str, list[float]] = {}
    subtotals: dict[int, float] = {}
    r = hrow + 1
    while r <= ws.max_row:
        name = _text(ws.cell(r, header.column).value)
        if not name:
            break
        values = [
            _number(ws.cell(r, c).value, f"{sheet}!{ws.cell(r, c).coordinate}")
            for c in year_cols.values()
        ]
        if name.lower().startswith("sub-total"):
            subtotals = dict(zip(year_cols, values))
            break
        families.append(name)
        rows[name] = values
        r += 1

    quantities = pd.DataFrame.from_dict(rows, orient="index", columns=list(year_cols))
    quantities.index.name = "Engine family"

    # Economics table: the second "Type" header on the same row.
    econ_header = _require(
        next(
            (
                ws.cell(hrow, c)
                for c in range(header.column + 1, ws.max_column + 1)
                if _text(ws.cell(hrow, c).value) == "Type"
            ),
            None,
        ),
        "the economics table header 'Type'",
        sheet,
    )
    headers = _header_columns(ws, hrow, econ_header.column + 1, ws.max_column)
    cols = {
        "income": _column_matching(headers, "income", sheet=sheet),
        "depreciation": _column_matching(headers, "depreciation", sheet=sheet),
        "interest": _column_matching(headers, "interest", sheet=sheet),
        "profit": _column_matching(headers, "profit", sheet=sheet),
    }
    econ: dict[str, dict[str, float]] = {}
    r = hrow + 1
    while r <= ws.max_row:
        name = _text(ws.cell(r, econ_header.column).value)
        if not name or name.lower().startswith("sub-total"):
            break
        econ[name] = {
            key: _number(ws.cell(r, c).value, f"{sheet}!{ws.cell(r, c).coordinate}")
            for key, c in cols.items()
        }
        r += 1
    economics = pd.DataFrame.from_dict(econ, orient="index")
    economics.index.name = "Engine family"

    economics_title = next(
        (
            _text(ws.cell(r, econ_header.column).value)
            for r in range(hrow - 1, 0, -1)
            if _text(ws.cell(r, econ_header.column).value)
        ),
        "Engine Type Economics",
    )
    return families, quantities, subtotals, economics, economics_title


def _parse_emissions(ws):
    sheet = ws.title
    header = _require(
        _find_cell(ws, lambda v: v == "Type"), "the emissions table header 'Type'", sheet
    )
    headers = _header_columns(ws, header.row, header.column + 1, ws.max_column)
    cols = {
        "per_flight": _column_matching(headers, "per engine", "flight", sheet=sheet),
        "flights": _column_matching(headers, "flights per year", sheet=sheet),
        "annual": _column_matching(headers, "annual", sheet=sheet),
    }
    rows: dict[str, dict[str, float]] = {}
    r = header.row + 1
    while r <= ws.max_row:
        name = _text(ws.cell(r, header.column).value)
        if not name:
            break
        rows[name] = {
            key: _number(ws.cell(r, c).value, f"{sheet}!{ws.cell(r, c).coordinate}")
            for key, c in cols.items()
        }
        r += 1
    emissions = pd.DataFrame.from_dict(rows, orient="index")
    emissions.index.name = "Engine family"
    return emissions


def _parse_projects(ws):
    sheet = ws.title
    header = _require(
        _find_cell(ws, lambda v: v == "Project Name"), "the 'Project Name' header", sheet
    )
    headers = _header_columns(ws, header.row, 1, ws.max_column)
    c_name = header.column
    c_link = _column_matching(headers, "project link", sheet=sheet)
    c_country = _column_matching(headers, "country", sheet=sheet)
    c_type = _column_matching(headers, "type", sheet=sheet)
    c_price = _column_matching(headers, "price", sheet=sheet)
    c_url = headers.get("Project URL")

    projects: list[Project] = []
    notes: list[str] = []
    url_mismatches: list[tuple[str, str, str]] = []
    for r in range(header.row + 1, ws.max_row + 1):
        name = _text(ws.cell(r, c_name).value)
        if not name:
            continue
        if name.lower().startswith("select from"):
            notes.append(
                f"Row {r} of '{sheet}' is a 'Select from drop-down' placeholder row, "
                "so it is not offered as a carbon project."
            )
            continue
        link_cell = ws.cell(r, c_link)
        link = link_cell.hyperlink.target if link_cell.hyperlink else None
        listed_url = _text(ws.cell(r, c_url).value) if c_url else ""
        if link and listed_url and listed_url != link:
            url_mismatches.append((name, link, listed_url))
        price = round(_number(ws.cell(r, c_price).value, f"{sheet}!R{r}"), 2)
        projects.append(
            Project(
                name=name,
                country=_text(ws.cell(r, c_country).value) or None,
                kind=_text(ws.cell(r, c_type).value),
                price=price,
                url=link or listed_url or None,
            )
        )

    if url_mismatches:
        name, link, listed = url_mismatches[0]
        notes.append(
            f"For {len(url_mismatches)} of {len(projects)} projects the 'Project URL' "
            "column points to a different page from the 'Project Link' hyperlink "
            f"(e.g. {name}: link → {link}; URL column → {listed}). The app uses the "
            "'Project Link' hyperlinks, whose addresses match the project names."
        )
    no_country = [p.name for p in projects if not p.country]
    if no_country:
        notes.append(
            "No country is given in the workbook for: " + ", ".join(no_country) + "."
        )
    return tuple(projects), notes


def load_workbook(path: Path | str = WORKBOOK_PATH) -> Workbook:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in (PORTFOLIO_SHEET, EMISSIONS_SHEET, CREDITS_SHEET):
        if sheet not in wb.sheetnames:
            raise WorkbookError(f"Workbook is missing the '{sheet}' sheet")

    families, quantities, subtotals, economics, econ_title = _parse_portfolio(
        wb[PORTFOLIO_SHEET]
    )
    emissions = _parse_emissions(wb[EMISSIONS_SHEET])
    projects, notes = _parse_projects(wb[CREDITS_SHEET])

    for label, table in (("Est.Emissions", emissions), ("economics", economics)):
        if set(table.index) != set(families):
            raise WorkbookError(
                f"Engine families in the {label} table do not match the portfolio"
            )
    emissions = emissions.loc[list(families)]
    economics = economics.loc[list(families)]
    if not projects:
        raise WorkbookError("No carbon projects found in the workbook")

    return Workbook(
        families=tuple(families),
        quantities=quantities,
        quantity_subtotals=subtotals,
        economics=economics,
        economics_title=econ_title,
        emissions=emissions,
        projects=projects,
        notes=tuple(notes),
        source_name=path.name,
    )
